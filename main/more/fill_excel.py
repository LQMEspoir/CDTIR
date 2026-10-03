#!/usr/bin/env python3
"""Fill backbone (full CDTIR) / few-shot / zero-shot / ablation experiment data into the experiment workbook."""
import glob, json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # add main/ to path for `from src...`

import openpyxl
from src.tourism_ie.evaluation.metrics import evaluate_records_detailed

ROOT = str(Path(__file__).resolve().parent.parent)
XLSX = str(Path(__file__).resolve().parent.parent / "experiment.xlsx")

MODEL_MAP = {
    "Qwen3-8B": "qwen3-8b",
    "Qwen2.5-7B-Instruct": "qwen2.5-7b",
    "GLM-4-9B-Chat": "glm4-9b",
    "Baichuan2-7B-Chat": "baichuan2-7b",
    "LLaMA-3.1-8B-Instruct": "llama3.1-8b",
    "Mistral-7B-Instruct-v0.3": "mistral7b-v0.3",
}


def load_metrics(pattern: str):
    fs = sorted(glob.glob(pattern))
    if not fs:
        return None
    d = json.load(open(fs[0]))
    return {
        "precision": d["micro"]["precision"],
        "recall": d["micro"]["recall"],
        "f1": d["micro"]["f1"],
        "accuracy": d.get("accuracy", 0.0),  # micro accuracy = TP/(TP+FP+FN)
    }


def fill_group(ws, start_row, dir_prefix):
    for i, (display, key) in enumerate(MODEL_MAP.items()):
        row = start_row + i
        base = f"{ROOT}/outputs/{dir_prefix}{key}"
        ner = load_metrics(f"{base}/ner_eval/*metrics.json")
        re_pred = load_metrics(f"{base}/re_eval/*metrics.json")
        re_gold = load_metrics(f"{base}/re_gold_eval/*metrics.json")
        ws.cell(row, 1, display.strip())
        for m, col in [("ner", 2), ("re", 6), ("re_gold", 10)]:
            val = {"ner": ner, "re": re_pred, "re_gold": re_gold}[m]
            if val:
                ws.cell(row, col, round(val["precision"], 6))
                ws.cell(row, col + 1, round(val["recall"], 6))
                ws.cell(row, col + 2, round(val["f1"], 6))
                ws.cell(row, col + 3, round(val["accuracy"], 6))


def _ablation_ner_accuracy(variant, b_ner_accuracy):
    """Ablation variants' NER micro accuracy: CDTIR/EH/TO reuse the backbone result; NMoRE/F are recomputed from predictions."""
    if variant in ("CDTIR", "CDTIR_EH", "CDTIR_TO"):
        return b_ner_accuracy
    rows = [json.loads(l) for l in open(f"{ROOT}/outputs/paper/{variant}/predictions/ner.jsonl") if l.strip()]
    return evaluate_records_detailed("ner", rows)["accuracy"]


def fill_ablation(ws):
    """Ablation sheet: 5 variants (CDTIR series), data from table_10.csv."""
    variants = ["CDTIR", "CDTIR_NMoRE", "CDTIR_F", "CDTIR_EH", "CDTIR_TO"]
    data = {}
    for l in open(f"{ROOT}/outputs/paper/reports/table_10.csv", encoding="utf-8-sig"):
        p = l.strip().split(",")
        if len(p) < 7 or p[0] not in variants:
            continue
        # fields: system,task,micro_p,micro_r,micro_f1,macro_f1,exact_accuracy,token_accuracy
        data[(p[0], p[1])] = {"p": float(p[2]), "r": float(p[3]), "f1": float(p[4])}
    # CDTIR full model = backbone, using the correct backbone_qwen2.5-7b result (avoid the ablation weighted drift)
    b_ner = load_metrics(f"{ROOT}/outputs/backbone_qwen2.5-7b/ner_eval/*metrics.json")
    b_rel = load_metrics(f"{ROOT}/outputs/backbone_qwen2.5-7b/re_eval/*metrics.json")
    for i, v in enumerate(variants):
        row = 3 + i
        ws.cell(row, 1, v)
        if v == "CDTIR":
            ner = {"p": b_ner["precision"], "r": b_ner["recall"], "f1": b_ner["f1"], "acc": b_ner["accuracy"]}
            rel = {"p": b_rel["precision"], "r": b_rel["recall"], "f1": b_rel["f1"]}
        elif v in ("CDTIR_EH", "CDTIR_TO"):
            # EH/TO reuse CDTIR's NER (base NER); RE uses their own variants
            ner = {"p": b_ner["precision"], "r": b_ner["recall"], "f1": b_ner["f1"], "acc": b_ner["accuracy"]}
            rel = data.get((v, "relation"))
        else:
            ner = data.get((v, "ner"))
            rel = data.get((v, "relation"))
            if ner:
                ner["acc"] = _ablation_ner_accuracy(v, b_ner["accuracy"])
        if ner:
            ws.cell(row, 2, round(ner["p"], 6))
            ws.cell(row, 3, round(ner["r"], 6))
            ws.cell(row, 4, round(ner["f1"], 6))
            ws.cell(row, 5, round(ner["acc"], 6))
        if rel:
            ws.cell(row, 6, round(rel["p"], 6))
            ws.cell(row, 7, round(rel["r"], 6))
            ws.cell(row, 8, round(rel["f1"], 6))
    # remove redundant model rows (original rows 8-9: Mistral)
    if ws.max_row > 7:
        ws.delete_rows(8, ws.max_row - 7)


wb = openpyxl.load_workbook(XLSX)
ws = wb["backbone"]

# restore group headers and clear header-row data columns (previously overwritten)
for r, title in [(3, "CDTIR (full)"), (12, "Few-Shot"), (21, "Zero-Shot")]:
    ws.cell(r, 1, title)
    for c in range(2, ws.max_column + 1):
        ws.cell(r, c).value = None

print("===== CDTIR full (backbone)=====")
fill_group(ws, 4, "backbone_")
print("===== Few-Shot =====")
fill_group(ws, 13, "few-shot/")
print("===== Zero-Shot =====")
fill_group(ws, 22, "zero-shot/")

print("===== ablation =====")
fill_ablation(wb["ablation"])

wb.save(XLSX)
print("\nsaved to", XLSX)
